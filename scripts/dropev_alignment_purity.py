#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Integrated pipeline script for DropEV-seq barcode-level alignment and purity analysis.

This script performs two main steps:
1. Align barcode-split paired-end reads to KP + EC combined reference genome
2. Execute reads number threshold sweep based on the filtered alignment results

Output organization:
1. `alignment/` subdirectory stores barcode alignment results and summary tables
2. `reads_num_threshold_sweep/` subdirectory stores threshold sweep summary results
3. Both steps share the same versioned output directory

Recommended execution:
    conda run -n evdna python scripts/dropev_alignment_purity.py \\
        --barcode-fastq-dir <path> \\
        --reference-fasta <path>
"""

from __future__ import annotations

import argparse
import csv
import multiprocessing as mp
import os
import re
import shutil
import subprocess
import sys
import traceback
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

import pandas as pd
from openpyxl import Workbook

try:
    import psutil
    import pysam
except ImportError as exc:  # pragma: no cover
    raise ImportError("[Error] Current evdna environment missing required dependencies. Please check psutil / pysam / openpyxl.") from exc


PROJECT_ROOT = Path(__file__).resolve().parents[1]
README_LOG = PROJECT_ROOT / "README_Log.md"
EXPECTED_CONDA_ENV = "evdna"
TASK_KEY = "barcode_reads_align_and_threshold_sweep"

DEFAULT_SAMPLE_ID = "d0"
DEFAULT_MAPQ_THRESHOLD = 60
DEFAULT_STRICT_MIN_IDENTITY = 0.95
DEFAULT_STRICT_MIN_ALIGNED_FRACTION = 0.80
DEFAULT_MIN_TOTAL_READS = 10
DEFAULT_SAVE_BARCODE_BAM = True
DEFAULT_PURITY_THRESHOLD_PERCENT = 85.0
DEFAULT_READS_THRESHOLDS = [200]

# Barcode filename prefix patterns
BARCODE_PREFIX_PATTERN_V1 = re.compile(r"^(?P<barcode_num>\d+)-(?P<barcode_core_id>[^-]+)$")
BARCODE_PREFIX_PATTERN_V2 = re.compile(
    r"^(?P<barcode_sample_id>.+)-(?P<barcode_num>\d+)-(?P<barcode_core_id>[^-]+)$"
)

ALIGNMENT_OUTPUT_COLUMNS = [
    "barcode_num",
    "barcode_id",
    "sample_prefix",
    "barcode_sample_id",
    "barcode_core_id",
    "primary_mapped_reads",
    "mapq_pass_primary_reads",
    "KP_numreads",
    "KP_covbases",
    "KP_coverage",
    "EC_numreads",
    "EC_covbases",
    "EC_coverage",
    "total_reads_uniq_paf",
    "KP_reads_paf",
    "EC_reads_paf",
    "ambiguous_reads_paf",
    "EC+KP_numreads",
    "passed_min_total_reads_filter",
    "barcode_result_dir",
    "barcode_summary_path",
    "barcode_bam_path",
]

# Global output paths (will be configured dynamically)
OUTPUT_DIR = PROJECT_ROOT / "output" / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{TASK_KEY}_v1"
ALIGNMENT_DIR = OUTPUT_DIR / "alignment"
ALIGNMENT_REPORTS_DIR = ALIGNMENT_DIR / "reports"
ALIGNMENT_INTERMEDIATE_DIR = ALIGNMENT_DIR / "intermediate"
BARCODE_RESULTS_DIR = ALIGNMENT_INTERMEDIATE_DIR / "barcode_alignment_results"
THRESHOLD_SWEEP_DIR = OUTPUT_DIR / "reads_num_threshold_sweep"
THRESHOLD_SWEEP_REPORTS_DIR = THRESHOLD_SWEEP_DIR / "reports"
SCRIPT_BACKUP_DIR = OUTPUT_DIR / "scripts_backup"

TASK_LOG_DIR = PROJECT_ROOT / "logs" / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{TASK_KEY}_v1"
PIPELINE_LOG_PATH = TASK_LOG_DIR / "barcode_reads_align_and_threshold_sweep.log"
ALIGNMENT_LOG_PATH = TASK_LOG_DIR / "alignment.log"
THRESHOLD_SWEEP_LOG_PATH = TASK_LOG_DIR / "reads_num_threshold_sweep.log"

ALIGNMENT_MAIN_TSV = ALIGNMENT_REPORTS_DIR / "d0_barcode_reads_alignment_summary_all_barcodes.tsv"
ALIGNMENT_MAIN_CSV = ALIGNMENT_REPORTS_DIR / "d0_barcode_reads_alignment_summary_all_barcodes.csv"
ALIGNMENT_MAIN_XLSX = ALIGNMENT_REPORTS_DIR / "d0_barcode_reads_alignment_summary_all_barcodes.xlsx"
ALIGNMENT_FILTERED_TSV = ALIGNMENT_REPORTS_DIR / "d0_barcode_reads_alignment_summary_filtered.tsv"
ALIGNMENT_FILTERED_CSV = ALIGNMENT_REPORTS_DIR / "d0_barcode_reads_alignment_summary_filtered.csv"
ALIGNMENT_FILTERED_XLSX = ALIGNMENT_REPORTS_DIR / "d0_barcode_reads_alignment_summary_filtered.xlsx"
ALIGNMENT_SUMMARY_TSV = ALIGNMENT_REPORTS_DIR / "d0_barcode_reads_alignment_run_summary.tsv"
ALIGNMENT_ERROR_TSV = ALIGNMENT_REPORTS_DIR / "d0_barcode_reads_alignment_errors.tsv"
ALIGNMENT_PROGRESS_TSV = ALIGNMENT_INTERMEDIATE_DIR / "d0_barcode_reads_alignment_progress.tsv"

THRESHOLD_SWEEP_SUMMARY_TSV = THRESHOLD_SWEEP_REPORTS_DIR / "d0_reads_num_threshold_sweep_summary.tsv"
THRESHOLD_SWEEP_SUMMARY_CSV = THRESHOLD_SWEEP_REPORTS_DIR / "d0_reads_num_threshold_sweep_summary.csv"
THRESHOLD_SWEEP_SUMMARY_XLSX = THRESHOLD_SWEEP_REPORTS_DIR / "d0_reads_num_threshold_sweep_summary.xlsx"
THRESHOLD_SWEEP_RUN_SUMMARY_TSV = THRESHOLD_SWEEP_REPORTS_DIR / "d0_reads_num_threshold_sweep_run_summary.tsv"
THRESHOLD_SWEEP_LEGEND_MD = THRESHOLD_SWEEP_REPORTS_DIR / "d0_reads_num_threshold_sweep_legend.md"

SINGLE_BARCODE_SUMMARY_FILENAME = "barcode_alignment_summary.tsv"
SINGLE_BARCODE_ERROR_FILENAME = "barcode_alignment_error.txt"


@dataclass
class ResourcePlan:
    """Resource planning results for alignment step."""

    cpu_count: int
    cpu_usage_percent: float
    total_memory_gb: float
    available_memory_gb: float
    available_memory_ratio: float
    worker_count: int
    bwa_threads_per_job: int


@dataclass
class BarcodeFastqPair:
    """Stores paired FASTQ paths and metadata for a single barcode."""

    barcode_num: int
    barcode_id: str
    sample_prefix: str
    barcode_sample_id: str
    barcode_core_id: str
    read1: str
    read2: str


@dataclass(frozen=True)
class ParsedBarcodePrefix:
    """Structured fields after parsing barcode filename prefix."""

    sample_prefix: str
    barcode_num: int | None
    barcode_id: str
    barcode_sample_id: str
    barcode_core_id: str


def print_and_log(message: str, log_handle) -> None:
    """Print message to both terminal and log file."""
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line = f"[{timestamp}] {message}"
    print(line, flush=True)
    log_handle.write(line + "\n")
    log_handle.flush()


def normalize_task_name(task_name: str) -> str:
    """Normalize task name to snake_case for version control and output naming."""
    normalized_name = re.sub(r"[^a-zA-Z0-9]+", "_", task_name.strip().lower()).strip("_")
    if not normalized_name:
        raise ValueError("[Error] Task name cannot be empty.")
    return normalized_name


def ensure_conda_environment() -> None:
    """Verify current execution is within evdna environment."""
    current_conda_env = os.environ.get("CONDA_DEFAULT_ENV", "").strip()
    if current_conda_env != EXPECTED_CONDA_ENV:
        raise EnvironmentError(
            f"[Error] Current Conda environment is '{current_conda_env or 'not detected'}'. "
            f"Please execute `conda activate {EXPECTED_CONDA_ENV}` or use "
            f"`conda run -n {EXPECTED_CONDA_ENV} python {Path(__file__).name}`."
        )


def ensure_required_tools() -> None:
    """Check if required external tools are available."""
    required_tools = ["bwa", "samtools"]
    missing_tools = [tool for tool in required_tools if shutil.which(tool) is None]
    if missing_tools:
        raise FileNotFoundError(f"[Error] Current environment missing required tools: {', '.join(missing_tools)}")


def parse_barcode_prefix(sample_prefix: str) -> ParsedBarcodePrefix:
    """
    Parse barcode filename prefix with backward compatibility.

    Supports two naming conventions:
    1. Version 1: `<cluster_size>-<barcode_seq>`
    2. Version 2: `<sample_id>-<cluster_size>-<barcode_seq>`

    Version 2 parsing uses right-to-left strategy to lock the last two fields,
    avoiding incorrect splitting when sample_id contains `-`.
    """
    v1_match = BARCODE_PREFIX_PATTERN_V1.fullmatch(sample_prefix)
    if v1_match is not None:
        barcode_num = int(v1_match.group("barcode_num"))
        barcode_core_id = v1_match.group("barcode_core_id")
        return ParsedBarcodePrefix(
            sample_prefix=sample_prefix,
            barcode_num=barcode_num,
            barcode_id=barcode_core_id,
            barcode_sample_id="",
            barcode_core_id=barcode_core_id,
        )

    v2_match = BARCODE_PREFIX_PATTERN_V2.fullmatch(sample_prefix)
    if v2_match is not None:
        barcode_sample_id = v2_match.group("barcode_sample_id")
        barcode_num = int(v2_match.group("barcode_num"))
        barcode_core_id = v2_match.group("barcode_core_id")
        return ParsedBarcodePrefix(
            sample_prefix=sample_prefix,
            barcode_num=barcode_num,
            barcode_id=f"{barcode_sample_id}-{barcode_core_id}",
            barcode_sample_id=barcode_sample_id,
            barcode_core_id=barcode_core_id,
        )

    return ParsedBarcodePrefix(
        sample_prefix=sample_prefix,
        barcode_num=None,
        barcode_id=sample_prefix,
        barcode_sample_id="",
        barcode_core_id=sample_prefix,
    )


def parse_args() -> argparse.Namespace:
    """Parse command-line arguments for integrated pipeline."""
    parser = argparse.ArgumentParser(description="Integrated barcode reads alignment and reads number threshold sweep.")
    parser.add_argument(
        "--barcode-fastq-dir",
        type=str,
        required=True,
        help="Directory containing barcode-split paired FASTQ files.",
    )
    parser.add_argument(
        "--reference-fasta",
        type=str,
        required=True,
        help="KP + EC combined reference FASTA file.",
    )
    parser.add_argument(
        "--sample-id",
        type=str,
        default=DEFAULT_SAMPLE_ID,
        help="Sample name for output file prefixes.",
    )
    parser.add_argument(
        "--mapq-threshold",
        type=int,
        default=DEFAULT_MAPQ_THRESHOLD,
        help="Minimum MAPQ threshold for strict filtering.",
    )
    parser.add_argument(
        "--strict-min-identity",
        type=float,
        default=DEFAULT_STRICT_MIN_IDENTITY,
        help="Minimum identity (0-1) for strict filtering.",
    )
    parser.add_argument(
        "--strict-min-aligned-fraction",
        type=float,
        default=DEFAULT_STRICT_MIN_ALIGNED_FRACTION,
        help="Minimum aligned fraction (0-1) for strict filtering.",
    )
    parser.add_argument(
        "--min-total-reads",
        type=int,
        default=DEFAULT_MIN_TOTAL_READS,
        help="Minimum total reads threshold for barcode filtering based on EC+KP_numreads.",
    )
    parser.add_argument(
        "--save-barcode-bam",
        action=argparse.BooleanOptionalAction,
        default=DEFAULT_SAVE_BARCODE_BAM,
        help="Whether to save sorted BAM and BAI for each barcode (default: enabled).",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Process only the first N barcodes for quick testing (default: process all).",
    )
    parser.add_argument(
        "--workers",
        type=int,
        default=None,
        help="Manually specify number of multiprocessing workers for alignment (default: auto-detect).",
    )
    parser.add_argument(
        "--bwa-threads-per-job",
        type=int,
        default=None,
        help="Number of bwa mem threads per barcode alignment job (default: 1).",
    )
    parser.add_argument(
        "--purity-threshold-percent",
        type=float,
        default=DEFAULT_PURITY_THRESHOLD_PERCENT,
        help="Purity threshold percentage for threshold sweep (e.g., 85 means 85%%).",
    )
    parser.add_argument(
        "--reads-thresholds",
        nargs="+",
        type=int,
        default=DEFAULT_READS_THRESHOLDS,
        help="List of reads number thresholds to sweep.",
    )
    args = parser.parse_args()
    if not (0.0 <= args.strict_min_identity <= 1.0):
        raise ValueError("[Error] --strict-min-identity must be between 0 and 1.")
    if not (0.0 <= args.strict_min_aligned_fraction <= 1.0):
        raise ValueError("[Error] --strict-min-aligned-fraction must be between 0 and 1.")
    return args


def detect_next_version(task_name: str) -> int:
    """Scan output directory and README_Log to find next version number, strictly avoiding rollback."""
    normalized_task_name = normalize_task_name(task_name)
    observed_versions: List[int] = []

    try:
        output_root = PROJECT_ROOT / "output"
        if output_root.exists():
            for directory_path in output_root.glob(f"*_{normalized_task_name}_v*"):
                if not directory_path.is_dir():
                    continue
                version_match = re.search(
                    rf"_{re.escape(normalized_task_name)}_v(\d+)$",
                    directory_path.name,
                )
                if version_match:
                    observed_versions.append(int(version_match.group(1)))
    except Exception as exc:
        raise RuntimeError(f"[Error] Failed to scan output directory for version numbers: {task_name}") from exc

    try:
        if README_LOG.exists():
            readme_text = README_LOG.read_text(encoding="utf-8")
            observed_versions.extend(
                int(version_text)
                for version_text in re.findall(
                    rf"{re.escape(normalized_task_name)}_v(\d+)",
                    readme_text,
                )
            )
    except Exception as exc:
        raise RuntimeError(f"[Error] Failed to read README_Log.md: {README_LOG}") from exc

    return max(observed_versions, default=0) + 1


def configure_output_paths(version_number: int, sample_id: str) -> None:
    """Configure main output directory and alignment / threshold sweep subdirectory filenames."""
    global OUTPUT_DIR
    global ALIGNMENT_DIR
    global ALIGNMENT_REPORTS_DIR
    global ALIGNMENT_INTERMEDIATE_DIR
    global BARCODE_RESULTS_DIR
    global THRESHOLD_SWEEP_DIR
    global THRESHOLD_SWEEP_REPORTS_DIR
    global SCRIPT_BACKUP_DIR
    global TASK_LOG_DIR
    global PIPELINE_LOG_PATH
    global ALIGNMENT_LOG_PATH
    global THRESHOLD_SWEEP_LOG_PATH
    global ALIGNMENT_MAIN_TSV
    global ALIGNMENT_MAIN_CSV
    global ALIGNMENT_MAIN_XLSX
    global ALIGNMENT_FILTERED_TSV
    global ALIGNMENT_FILTERED_CSV
    global ALIGNMENT_FILTERED_XLSX
    global ALIGNMENT_SUMMARY_TSV
    global ALIGNMENT_ERROR_TSV
    global ALIGNMENT_PROGRESS_TSV
    global THRESHOLD_SWEEP_SUMMARY_TSV
    global THRESHOLD_SWEEP_SUMMARY_CSV
    global THRESHOLD_SWEEP_SUMMARY_XLSX
    global THRESHOLD_SWEEP_RUN_SUMMARY_TSV
    global THRESHOLD_SWEEP_LEGEND_MD

    timestamp_text = datetime.now().strftime("%Y%m%d_%H%M%S")
    normalized_sample_id = normalize_task_name(sample_id)

    OUTPUT_DIR = PROJECT_ROOT / "output" / f"{timestamp_text}_{TASK_KEY}_v{version_number}"
    ALIGNMENT_DIR = OUTPUT_DIR / "alignment"
    ALIGNMENT_REPORTS_DIR = ALIGNMENT_DIR / "reports"
    ALIGNMENT_INTERMEDIATE_DIR = ALIGNMENT_DIR / "intermediate"
    BARCODE_RESULTS_DIR = ALIGNMENT_INTERMEDIATE_DIR / "barcode_alignment_results"
    THRESHOLD_SWEEP_DIR = OUTPUT_DIR / "reads_num_threshold_sweep"
    THRESHOLD_SWEEP_REPORTS_DIR = THRESHOLD_SWEEP_DIR / "reports"
    SCRIPT_BACKUP_DIR = OUTPUT_DIR / "scripts_backup"

    TASK_LOG_DIR = PROJECT_ROOT / "logs" / f"{timestamp_text}_{TASK_KEY}_v{version_number}"
    PIPELINE_LOG_PATH = TASK_LOG_DIR / "barcode_reads_align_and_threshold_sweep.log"
    ALIGNMENT_LOG_PATH = TASK_LOG_DIR / "alignment.log"
    THRESHOLD_SWEEP_LOG_PATH = TASK_LOG_DIR / "reads_num_threshold_sweep.log"

    ALIGNMENT_MAIN_TSV = ALIGNMENT_REPORTS_DIR / f"{normalized_sample_id}_barcode_reads_alignment_summary_all_barcodes.tsv"
    ALIGNMENT_MAIN_CSV = ALIGNMENT_REPORTS_DIR / f"{normalized_sample_id}_barcode_reads_alignment_summary_all_barcodes.csv"
    ALIGNMENT_MAIN_XLSX = ALIGNMENT_REPORTS_DIR / f"{normalized_sample_id}_barcode_reads_alignment_summary_all_barcodes.xlsx"
    ALIGNMENT_FILTERED_TSV = ALIGNMENT_REPORTS_DIR / f"{normalized_sample_id}_barcode_reads_alignment_summary_filtered.tsv"
    ALIGNMENT_FILTERED_CSV = ALIGNMENT_REPORTS_DIR / f"{normalized_sample_id}_barcode_reads_alignment_summary_filtered.csv"
    ALIGNMENT_FILTERED_XLSX = ALIGNMENT_REPORTS_DIR / f"{normalized_sample_id}_barcode_reads_alignment_summary_filtered.xlsx"
    ALIGNMENT_SUMMARY_TSV = ALIGNMENT_REPORTS_DIR / f"{normalized_sample_id}_barcode_reads_alignment_run_summary.tsv"
    ALIGNMENT_ERROR_TSV = ALIGNMENT_REPORTS_DIR / f"{normalized_sample_id}_barcode_reads_alignment_errors.tsv"
    ALIGNMENT_PROGRESS_TSV = ALIGNMENT_INTERMEDIATE_DIR / f"{normalized_sample_id}_barcode_reads_alignment_progress.tsv"

    THRESHOLD_SWEEP_SUMMARY_TSV = THRESHOLD_SWEEP_REPORTS_DIR / f"{normalized_sample_id}_reads_num_threshold_sweep_summary.tsv"
    THRESHOLD_SWEEP_SUMMARY_CSV = THRESHOLD_SWEEP_REPORTS_DIR / f"{normalized_sample_id}_reads_num_threshold_sweep_summary.csv"
    THRESHOLD_SWEEP_SUMMARY_XLSX = THRESHOLD_SWEEP_REPORTS_DIR / f"{normalized_sample_id}_reads_num_threshold_sweep_summary.xlsx"
    THRESHOLD_SWEEP_RUN_SUMMARY_TSV = THRESHOLD_SWEEP_REPORTS_DIR / f"{normalized_sample_id}_reads_num_threshold_sweep_run_summary.tsv"
    THRESHOLD_SWEEP_LEGEND_MD = THRESHOLD_SWEEP_REPORTS_DIR / f"{normalized_sample_id}_reads_num_threshold_sweep_legend.md"


def ensure_reference_indexes(reference_fasta: Path) -> None:
    """Check if bwa index files are complete."""
    required_suffixes = [".amb", ".ann", ".bwt", ".pac", ".sa"]
    missing_index_files = [reference_fasta.with_suffix(reference_fasta.suffix + suffix) for suffix in required_suffixes]
    missing_index_files = [path for path in missing_index_files if not path.exists()]
    if missing_index_files:
        missing_text = ", ".join(str(path) for path in missing_index_files)
        raise FileNotFoundError(f"[Error] Missing bwa index files: {missing_text}")


def ensure_master_paths(barcode_fastq_dir: Path, reference_fasta: Path) -> None:
    """Check input paths and create main output directory and subdirectories."""
    if not barcode_fastq_dir.exists():
        raise FileNotFoundError(f"[Error] Barcode FASTQ directory not found: {barcode_fastq_dir}")
    if not reference_fasta.exists():
        raise FileNotFoundError(f"[Error] Combined reference not found: {reference_fasta}")

    reference_fai = reference_fasta.with_suffix(reference_fasta.suffix + ".fai")
    if not reference_fai.exists():
        raise FileNotFoundError(f"[Error] Reference index file not found: {reference_fai}")
    ensure_reference_indexes(reference_fasta)

    for directory in [
        OUTPUT_DIR,
        ALIGNMENT_DIR,
        ALIGNMENT_REPORTS_DIR,
        ALIGNMENT_INTERMEDIATE_DIR,
        BARCODE_RESULTS_DIR,
        THRESHOLD_SWEEP_DIR,
        THRESHOLD_SWEEP_REPORTS_DIR,
        SCRIPT_BACKUP_DIR,
        TASK_LOG_DIR,
    ]:
        directory.mkdir(parents=True, exist_ok=True)


def detect_alignment_resources(requested_workers: int | None = None, requested_bwa_threads: int | None = None) -> ResourcePlan:
    """Plan alignment parallelism based on current CPU / memory load."""
    cpu_count = psutil.cpu_count(logical=True) or 1
    cpu_usage_percent = psutil.cpu_percent(interval=1)
    memory_info = psutil.virtual_memory()
    total_memory_gb = memory_info.total / (1024**3)
    available_memory_gb = memory_info.available / (1024**3)
    available_memory_ratio = memory_info.available / memory_info.total

    if requested_workers is not None:
        worker_count = max(1, requested_workers)
    elif cpu_usage_percent <= 75 and available_memory_ratio >= 0.25:
        worker_count = min(max(cpu_count // 8, 8), 24)
    else:
        worker_count = 4

    if requested_bwa_threads is not None:
        bwa_threads_per_job = max(1, requested_bwa_threads)
    else:
        bwa_threads_per_job = 1

    return ResourcePlan(
        cpu_count=cpu_count,
        cpu_usage_percent=cpu_usage_percent,
        total_memory_gb=total_memory_gb,
        available_memory_gb=available_memory_gb,
        available_memory_ratio=available_memory_ratio,
        worker_count=worker_count,
        bwa_threads_per_job=bwa_threads_per_job,
    )


def load_reference_lengths(reference_fai: Path) -> Dict[str, int]:
    """Read .fai file and calculate total reference length for KP / EC."""
    kp_total_bases = 0
    ec_total_bases = 0

    try:
        with reference_fai.open("r", encoding="utf-8") as handle:
            for line in handle:
                parts = line.rstrip("\n").split("\t")
                if len(parts) < 2:
                    continue
                contig_name = parts[0]
                contig_length = int(parts[1])
                if contig_name.startswith("kp_"):
                    kp_total_bases += contig_length
                elif contig_name.startswith("ec_"):
                    ec_total_bases += contig_length
    except Exception as exc:
        raise RuntimeError(f"[Error] Failed to read reference fai file: {reference_fai}") from exc

    return {
        "KP_total_bases": kp_total_bases,
        "EC_total_bases": ec_total_bases,
    }


def collect_barcode_pairs(barcode_fastq_dir: Path, limit: int | None = None) -> List[BarcodeFastqPair]:
    """Scan directory and collect paired barcode FASTQ files."""
    read1_files = sorted(barcode_fastq_dir.glob("*_S1_L001_R1_001.fastq.gz"))
    pair_records: List[BarcodeFastqPair] = []

    for read1_path in read1_files:
        read2_path = barcode_fastq_dir / read1_path.name.replace("_R1_001.fastq.gz", "_R2_001.fastq.gz")
        if not read2_path.exists():
            print(f"[Warning] Missing paired R2, skipped: {read1_path}", file=sys.stderr)
            continue

        sample_prefix = read1_path.name.replace("_S1_L001_R1_001.fastq.gz", "")
        parsed_prefix = parse_barcode_prefix(sample_prefix)
        if parsed_prefix.barcode_num is None:
            print(f"[Warning] Unable to parse barcode prefix, skipped: {sample_prefix}", file=sys.stderr)
            continue

        pair_records.append(
            BarcodeFastqPair(
                barcode_num=parsed_prefix.barcode_num,
                barcode_id=parsed_prefix.barcode_id,
                sample_prefix=parsed_prefix.sample_prefix,
                barcode_sample_id=parsed_prefix.barcode_sample_id,
                barcode_core_id=parsed_prefix.barcode_core_id,
                read1=str(read1_path),
                read2=str(read2_path),
            )
        )

    pair_records.sort(key=lambda record: record.barcode_num, reverse=True)
    if limit is not None:
        pair_records = pair_records[:limit]
    return pair_records


def merge_intervals(intervals: List[Tuple[int, int]]) -> int:
    """Calculate union length of half-open interval list."""
    if not intervals:
        return 0

    sorted_intervals = sorted(intervals)
    merged_start, merged_end = sorted_intervals[0]
    covered_bases = 0

    for current_start, current_end in sorted_intervals[1:]:
        if current_start <= merged_end:
            merged_end = max(merged_end, current_end)
        else:
            covered_bases += merged_end - merged_start
            merged_start, merged_end = current_start, current_end

    covered_bases += merged_end - merged_start
    return covered_bases


def infer_species_from_reference(reference_name: str | None) -> str | None:
    """Identify species based on reference sequence name prefix."""
    if reference_name is None:
        return None
    if reference_name.startswith("kp_"):
        return "KP"
    if reference_name.startswith("ec_"):
        return "EC"
    return None


def compute_alignment_identity(alignment: pysam.AlignedSegment) -> float:
    """Estimate alignment identity using NM tag."""
    aligned_length = alignment.query_alignment_length or 0
    if aligned_length <= 0:
        return 0.0
    if alignment.has_tag("NM"):
        edit_distance = int(alignment.get_tag("NM"))
        return max(0.0, 1.0 - (edit_distance / aligned_length))
    return 0.0


def compute_alignment_fraction(alignment: pysam.AlignedSegment) -> float:
    """Estimate alignment coverage fraction on query."""
    query_length = alignment.query_length or alignment.infer_read_length() or 0
    aligned_length = alignment.query_alignment_length or 0
    if query_length <= 0:
        return 0.0
    return aligned_length / query_length


def sort_and_index_bam(unsorted_bam_path: Path, sorted_bam_path: Path) -> None:
    """Sort temporary BAM and build index for per-barcode review."""
    try:
        pysam.sort("-o", str(sorted_bam_path), str(unsorted_bam_path))
        pysam.index(str(sorted_bam_path))
    except Exception as exc:
        raise RuntimeError(f"[Error] Failed to sort or index BAM: {unsorted_bam_path}") from exc
    finally:
        if unsorted_bam_path.exists():
            unsorted_bam_path.unlink()


def write_table_with_columns(rows: Iterable[Dict[str, object]], output_path: Path, columns: Sequence[str], delimiter: str) -> None:
    """Write table with specified column order."""
    with output_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(columns), delimiter=delimiter)
        writer.writeheader()
        for row in rows:
            formatted_row = dict(row)
            for coverage_key in ["KP_coverage", "EC_coverage"]:
                formatted_row[coverage_key] = f"{float(row[coverage_key]):.6f}"
            writer.writerow(formatted_row)


def write_single_barcode_summary(output_path: Path, row: Dict[str, object]) -> None:
    """Write single barcode statistics to its independent result directory."""
    write_table_with_columns([row], output_path, ALIGNMENT_OUTPUT_COLUMNS, "\t")


def process_single_barcode(
    task_input: Tuple[BarcodeFastqPair, str, int, int, int, int, float, float, int, str, bool]
) -> Dict[str, object]:
    """Process single barcode alignment, BAM output, and statistics."""
    (
        barcode_record,
        reference_fasta,
        bwa_threads,
        kp_total_bases,
        ec_total_bases,
        mapq_threshold,
        strict_min_identity,
        strict_min_aligned_fraction,
        min_total_reads,
        barcode_results_dir,
        save_barcode_bam,
    ) = task_input

    process = None
    alignment_file = None
    bam_writer = None
    unsorted_bam_path = None
    sorted_bam_path = None
    barcode_output_dir = Path(barcode_results_dir) / barcode_record.sample_prefix
    single_barcode_summary_path = barcode_output_dir / SINGLE_BARCODE_SUMMARY_FILENAME

    try:
        barcode_output_dir.mkdir(parents=True, exist_ok=True)
        command = [
            "bwa",
            "mem",
            "-t",
            str(bwa_threads),
            reference_fasta,
            barcode_record.read1,
            barcode_record.read2,
        ]

        process = subprocess.Popen(
            command,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=False,
        )
        if process.stdout is None:
            raise RuntimeError("[Error] bwa mem did not return stdout handle.")

        alignment_file = pysam.AlignmentFile(process.stdout, "r")
        if save_barcode_bam:
            unsorted_bam_path = barcode_output_dir / f"{barcode_record.sample_prefix}.primary_mapped.unsorted.bam"
            sorted_bam_path = barcode_output_dir / f"{barcode_record.sample_prefix}.primary_mapped.sorted.bam"
            bam_writer = pysam.AlignmentFile(str(unsorted_bam_path), "wb", template=alignment_file)
        else:
            sorted_bam_path = None

        all_query_names: set[str] = set()
        strict_species_hits: Dict[str, set[str]] = {}
        strict_intervals_by_query_species: Dict[str, Dict[str, Dict[str, List[Tuple[int, int]]]]] = {}
        primary_mapped_reads = 0
        mapq_pass_primary_reads = 0

        for alignment in alignment_file:
            if alignment.is_unmapped or alignment.is_secondary or alignment.is_supplementary:
                continue
            species = infer_species_from_reference(alignment.reference_name)
            query_name = alignment.query_name
            if species is None or query_name is None:
                continue

            primary_mapped_reads += 1
            all_query_names.add(query_name)
            if bam_writer is not None:
                bam_writer.write(alignment)

            identity = compute_alignment_identity(alignment)
            aligned_fraction = compute_alignment_fraction(alignment)
            if alignment.mapping_quality < mapq_threshold:
                continue
            if identity < strict_min_identity:
                continue
            if aligned_fraction < strict_min_aligned_fraction:
                continue

            reference_name = alignment.reference_name
            start = int(alignment.reference_start)
            end = int(alignment.reference_end or alignment.reference_start)
            if end <= start:
                continue

            mapq_pass_primary_reads += 1
            strict_species_hits.setdefault(query_name, set()).add(species)
            strict_intervals_by_query_species.setdefault(query_name, {}).setdefault(species, {}).setdefault(reference_name, []).append((start, end))

        if bam_writer is not None:
            bam_writer.close()
            bam_writer = None

        if alignment_file is not None:
            alignment_file.close()
            alignment_file = None

        stderr_output = process.stderr.read().decode("utf-8", errors="replace") if process.stderr is not None else ""
        return_code = process.wait()
        if return_code != 0:
            raise RuntimeError(
                f"[Error] bwa mem processing barcode {barcode_record.sample_prefix} failed with return code {return_code}."
                f" stderr: {stderr_output[:500]}"
            )

        if save_barcode_bam and unsorted_bam_path is not None and sorted_bam_path is not None:
            sort_and_index_bam(unsorted_bam_path, sorted_bam_path)

        kp_intervals: Dict[str, List[Tuple[int, int]]] = {}
        ec_intervals: Dict[str, List[Tuple[int, int]]] = {}
        strict_kp_count = 0
        strict_ec_count = 0
        strict_ambiguous_count = 0

        for query_name in all_query_names:
            species_set = strict_species_hits.get(query_name, set())
            if species_set == {"KP"}:
                strict_kp_count += 1
                for reference_name, intervals in strict_intervals_by_query_species.get(query_name, {}).get("KP", {}).items():
                    kp_intervals.setdefault(reference_name, []).extend(intervals)
            elif species_set == {"EC"}:
                strict_ec_count += 1
                for reference_name, intervals in strict_intervals_by_query_species.get(query_name, {}).get("EC", {}).items():
                    ec_intervals.setdefault(reference_name, []).extend(intervals)
            elif species_set:
                strict_ambiguous_count += 1

        kp_covbases = sum(merge_intervals(intervals) for intervals in kp_intervals.values())
        ec_covbases = sum(merge_intervals(intervals) for intervals in ec_intervals.values())
        strict_informative_reads = strict_kp_count + strict_ec_count

        result_row = {
            "barcode_num": barcode_record.barcode_num,
            "barcode_id": barcode_record.barcode_id,
            "sample_prefix": barcode_record.sample_prefix,
            "barcode_sample_id": barcode_record.barcode_sample_id,
            "barcode_core_id": barcode_record.barcode_core_id,
            "primary_mapped_reads": primary_mapped_reads,
            "mapq_pass_primary_reads": mapq_pass_primary_reads,
            "KP_numreads": strict_kp_count,
            "KP_covbases": kp_covbases,
            "KP_coverage": round((kp_covbases / kp_total_bases) if kp_total_bases else 0.0, 6),
            "EC_numreads": strict_ec_count,
            "EC_covbases": ec_covbases,
            "EC_coverage": round((ec_covbases / ec_total_bases) if ec_total_bases else 0.0, 6),
            "total_reads_uniq_paf": strict_informative_reads + strict_ambiguous_count,
            "KP_reads_paf": strict_kp_count,
            "EC_reads_paf": strict_ec_count,
            "ambiguous_reads_paf": strict_ambiguous_count,
            "EC+KP_numreads": strict_informative_reads,
            "passed_min_total_reads_filter": "yes" if strict_informative_reads >= min_total_reads else "no",
            "barcode_result_dir": str(barcode_output_dir),
            "barcode_summary_path": str(single_barcode_summary_path),
            "barcode_bam_path": str(sorted_bam_path) if sorted_bam_path is not None else "",
        }
        write_single_barcode_summary(single_barcode_summary_path, result_row)
        return {
            "status": "ok",
            "sample_prefix": barcode_record.sample_prefix,
            "result": result_row,
        }
    except Exception as exc:
        try:
            barcode_output_dir.mkdir(parents=True, exist_ok=True)
            (barcode_output_dir / SINGLE_BARCODE_ERROR_FILENAME).write_text(
                f"{str(exc)}\n\n{traceback.format_exc()}",
                encoding="utf-8",
            )
        except Exception:
            pass
        return {
            "status": "error",
            "sample_prefix": barcode_record.sample_prefix,
            "error_message": str(exc),
            "traceback": traceback.format_exc(),
        }
    finally:
        try:
            if bam_writer is not None:
                bam_writer.close()
        except Exception:
            pass
        try:
            if alignment_file is not None:
                alignment_file.close()
        except Exception:
            pass
        try:
            if process is not None:
                if process.stdout is not None:
                    process.stdout.close()
                if process.stderr is not None:
                    process.stderr.close()
        except Exception:
            pass


def write_alignment_xlsx(rows: List[Dict[str, object]], output_path: Path) -> None:
    """Write alignment Excel summary table."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "barcode_reads"
    worksheet.append(ALIGNMENT_OUTPUT_COLUMNS)

    for row in rows:
        worksheet.append([row[column] for column in ALIGNMENT_OUTPUT_COLUMNS])

    for column_name in ["KP_coverage", "EC_coverage"]:
        column_index = ALIGNMENT_OUTPUT_COLUMNS.index(column_name) + 1
        for column_cells in worksheet.iter_cols(
            min_col=column_index,
            max_col=column_index,
            min_row=2,
            max_row=worksheet.max_row,
        ):
            for cell in column_cells:
                cell.number_format = "0.000000"

    workbook.save(output_path)


def write_alignment_summary(
    barcode_fastq_dir: Path,
    reference_fasta: Path,
    sample_id: str,
    total_barcodes: int,
    success_count: int,
    filtered_count: int,
    error_count: int,
    reference_lengths: Dict[str, int],
    resource_plan: ResourcePlan,
    rows_all: List[Dict[str, object]],
    rows_filtered: List[Dict[str, object]],
    mapq_threshold: int,
    strict_min_identity: float,
    strict_min_aligned_fraction: float,
    min_total_reads: int,
    save_barcode_bam: bool,
) -> None:
    """Write sample-level alignment summary."""
    summary_rows = [
        ("sample_id", sample_id),
        ("barcode_fastq_dir", str(barcode_fastq_dir)),
        ("reference_fasta", str(reference_fasta)),
        ("total_barcodes", total_barcodes),
        ("success_count", success_count),
        ("filtered_count", filtered_count),
        ("error_count", error_count),
        ("mapq_threshold", mapq_threshold),
        ("strict_min_identity", f"{strict_min_identity:.3f}"),
        ("strict_min_aligned_fraction", f"{strict_min_aligned_fraction:.3f}"),
        ("min_total_reads", min_total_reads),
        ("save_barcode_bam", str(save_barcode_bam)),
        ("KP_total_reference_bases", reference_lengths["KP_total_bases"]),
        ("EC_total_reference_bases", reference_lengths["EC_total_bases"]),
        ("summed_primary_mapped_reads", sum(int(row["primary_mapped_reads"]) for row in rows_all)),
        ("summed_mapq_pass_primary_reads", sum(int(row["mapq_pass_primary_reads"]) for row in rows_all)),
        ("summed_KP_numreads_all", sum(int(row["KP_numreads"]) for row in rows_all)),
        ("summed_EC_numreads_all", sum(int(row["EC_numreads"]) for row in rows_all)),
        ("summed_KP_numreads_filtered", sum(int(row["KP_numreads"]) for row in rows_filtered)),
        ("summed_EC_numreads_filtered", sum(int(row["EC_numreads"]) for row in rows_filtered)),
        ("summed_KP_covbases_filtered", sum(int(row["KP_covbases"]) for row in rows_filtered)),
        ("summed_EC_covbases_filtered", sum(int(row["EC_covbases"]) for row in rows_filtered)),
        ("summed_KP_reads_paf_unique_only_filtered", sum(int(row["KP_reads_paf"]) for row in rows_filtered)),
        ("summed_EC_reads_paf_unique_only_filtered", sum(int(row["EC_reads_paf"]) for row in rows_filtered)),
        ("summed_ambiguous_reads_paf_filtered", sum(int(row["ambiguous_reads_paf"]) for row in rows_filtered)),
        ("cpu_count", resource_plan.cpu_count),
        ("cpu_usage_percent", f"{resource_plan.cpu_usage_percent:.2f}"),
        ("total_memory_gb", f"{resource_plan.total_memory_gb:.2f}"),
        ("available_memory_gb", f"{resource_plan.available_memory_gb:.2f}"),
        ("worker_count", resource_plan.worker_count),
        ("bwa_threads_per_job", resource_plan.bwa_threads_per_job),
    ]

    with ALIGNMENT_SUMMARY_TSV.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["metric", "value"])
        writer.writerows(summary_rows)


def validate_reads_thresholds(reads_thresholds: List[int]) -> List[int]:
    """Clean and sort threshold sweep threshold list."""
    cleaned_thresholds = sorted({int(threshold) for threshold in reads_thresholds})
    if not cleaned_thresholds:
        raise ValueError("[Error] Reads threshold list cannot be empty.")
    if any(threshold < 0 for threshold in cleaned_thresholds):
        raise ValueError("[Error] Reads thresholds cannot be negative.")
    return cleaned_thresholds


def load_threshold_sweep_input_dataframe(input_file: Path) -> pd.DataFrame:
    """Read filtered TSV, extract strict informative KP/EC reads and construct Total_reads."""
    try:
        df = pd.read_csv(input_file, sep="\t")
    except Exception as exc:
        raise RuntimeError(f"[Error] Failed to read threshold sweep input file: {input_file}") from exc

    required_cols = {"barcode_num", "barcode_id", "KP_reads_paf", "EC_reads_paf"}
    missing = required_cols - set(df.columns)
    if missing:
        raise ValueError(f"[Error] Threshold sweep input file missing required columns: {sorted(missing)}")

    df["KP_reads"] = pd.to_numeric(df["KP_reads_paf"], errors="coerce").fillna(0).astype(int)
    df["EC_reads"] = pd.to_numeric(df["EC_reads_paf"], errors="coerce").fillna(0).astype(int)
    df["Total_reads"] = df["KP_reads"] + df["EC_reads"]
    return df


def classify_barcode(ratio_kp: float, purity_threshold: float) -> str:
    """Classify barcode based on KP ratio and purity threshold."""
    if ratio_kp >= purity_threshold:
        return "Pure KP"
    if ratio_kp <= (1 - purity_threshold):
        return "Pure EC"
    return "Mixed"


def sweep_thresholds(df: pd.DataFrame, reads_thresholds: List[int], purity_threshold: float) -> pd.DataFrame:
    """Execute filtering and classification statistics for each reads threshold."""
    results = []

    for threshold in reads_thresholds:
        filtered = df[df["Total_reads"] >= threshold].copy()
        total_valid = len(filtered)

        if total_valid == 0:
            results.append(
                {
                    "min_reads_threshold": threshold,
                    "purity_threshold": f"{purity_threshold * 100:.0f}%",
                    "valid_barcodes": 0,
                    "Pure_KP": 0,
                    "Pure_EC": 0,
                    "Mixed": 0,
                    "Pure_KP_pct": "0.0%",
                    "Pure_EC_pct": "0.0%",
                    "Mixed_pct": "0.0%",
                    "doublet_rate": "0.0%",
                }
            )
            continue

        filtered["Ratio_KP"] = filtered["KP_reads"] / filtered["Total_reads"]
        filtered["Category"] = filtered["Ratio_KP"].apply(lambda ratio: classify_barcode(ratio, purity_threshold))

        counts = filtered["Category"].value_counts()
        n_kp = int(counts.get("Pure KP", 0))
        n_ec = int(counts.get("Pure EC", 0))
        n_mixed = int(counts.get("Mixed", 0))

        results.append(
            {
                "min_reads_threshold": threshold,
                "purity_threshold": f"{purity_threshold * 100:.0f}%",
                "valid_barcodes": total_valid,
                "Pure_KP": n_kp,
                "Pure_EC": n_ec,
                "Mixed": n_mixed,
                "Pure_KP_pct": f"{n_kp / total_valid * 100:.1f}%",
                "Pure_EC_pct": f"{n_ec / total_valid * 100:.1f}%",
                "Mixed_pct": f"{n_mixed / total_valid * 100:.1f}%",
                "doublet_rate": f"{n_mixed / total_valid * 100:.1f}%",
            }
        )

    return pd.DataFrame(results)


def write_threshold_sweep_run_summary(
    input_file: Path,
    sample_id: str,
    purity_threshold_percent: float,
    reads_thresholds: List[int],
    input_barcode_count: int,
    resource_status: Dict[str, float],
) -> None:
    """Write threshold sweep step run summary."""
    summary_rows = [
        ("sample_id", sample_id),
        ("input_file", str(input_file)),
        ("input_barcode_count", input_barcode_count),
        ("purity_threshold_percent", f"{purity_threshold_percent:.2f}"),
        ("reads_thresholds", ",".join(str(value) for value in reads_thresholds)),
        ("threshold_count", len(reads_thresholds)),
        ("cpu_count", int(resource_status["cpu_count"])),
        ("cpu_usage_percent", f"{resource_status['cpu_usage_percent']:.2f}"),
        ("available_memory_gb", f"{resource_status['available_memory_gb']:.2f}"),
    ]

    run_summary_df = pd.DataFrame(summary_rows, columns=["metric", "value"])
    run_summary_df.to_csv(THRESHOLD_SWEEP_RUN_SUMMARY_TSV, sep="\t", index=False)


def write_threshold_sweep_legend(
    input_file: Path,
    input_barcode_count: int,
    purity_threshold_percent: float,
    reads_thresholds: List[int],
) -> None:
    """Write threshold sweep legend description."""
    legend_text = f"""# Reads Num Threshold Sweep Summary Legend

This table summarizes classification results under different minimum informative reads thresholds
based on DropEV-seq dual-species Mock sample barcode reads alignment results.

Description:
- Data source: `{input_file}`, total {input_barcode_count} barcodes.
- Input criteria: Filtered TSV output from alignment step of this integrated pipeline.
- `KP_reads_paf` / `EC_reads_paf` come from strict filtering:
  First filter by `MAPQ`, `identity`, `aligned_fraction`, then classify by query_name into KP-only / EC-only / ambiguous.
- Valid barcode definition: `KP_reads_paf + EC_reads_paf >= minimum informative reads threshold`.
- Purity threshold fixed at {purity_threshold_percent:.0f}%:
  KP ratio >= {purity_threshold_percent:.0f}% classified as Pure KP;
  EC ratio >= {purity_threshold_percent:.0f}% classified as Pure EC;
  Others classified as Mixed.
- Scanned threshold list: {", ".join(str(value) for value in reads_thresholds)}.
- `doublet_rate` defined as Mixed ratio among valid barcodes, useful for evaluating cross-droplet contamination.
"""
    THRESHOLD_SWEEP_LEGEND_MD.write_text(legend_text, encoding="utf-8")


def backup_current_script() -> Path:
    """Backup current integrated script to output directory."""
    target_path = SCRIPT_BACKUP_DIR / Path(__file__).name
    shutil.copy2(Path(__file__), target_path)
    return target_path


def append_readme_log(
    output_dir: Path,
    sample_id: str,
    barcode_fastq_dir: Path,
    reference_fasta: Path,
    mapq_threshold: int,
    strict_min_identity: float,
    strict_min_aligned_fraction: float,
    min_total_reads: int,
    purity_threshold_percent: float,
    reads_thresholds: List[int],
) -> None:
    """Record this integrated pipeline run to README_Log.md."""
    today_text = datetime.now().strftime("%Y-%m-%d")
    threshold_text = ",".join(str(value) for value in reads_thresholds)
    log_block = "\n".join(
        [
            f"- **Date**: {today_text}",
            f"- **Script/Task**: `scripts/{Path(__file__).name}`",
            "- **Core Function**: [Integrated execution of barcode reads alignment to KP+EC combined reference with strict filtering, and reads number threshold sweep based on filtered TSV informative reads]",
            f"- **Key Parameters/Results**: [sample_id={sample_id}; barcode_fastq_dir={barcode_fastq_dir}; reference={reference_fasta}; mapq_threshold={mapq_threshold}; strict_min_identity={strict_min_identity:.2f}; strict_min_aligned_fraction={strict_min_aligned_fraction:.2f}; min_total_reads={min_total_reads}; purity_threshold_percent={purity_threshold_percent:.0f}; reads_thresholds={threshold_text}]",
            f"- **Output Path**: `output/{output_dir.name}/`",
            "",
        ]
    )

    try:
        existing_text = README_LOG.read_text(encoding="utf-8") if README_LOG.exists() else "# README Log\n\n"
        README_LOG.write_text(log_block + existing_text, encoding="utf-8")
    except Exception as exc:
        raise RuntimeError(f"[Error] Failed to write README_Log.md: {README_LOG}") from exc


def run_alignment_step(
    barcode_fastq_dir: Path,
    reference_fasta: Path,
    sample_id: str,
    mapq_threshold: int,
    strict_min_identity: float,
    strict_min_aligned_fraction: float,
    min_total_reads: int,
    save_barcode_bam: bool,
    limit: int | None,
    workers: int | None,
    bwa_threads_per_job: int | None,
    pipeline_log_handle,
) -> Dict[str, object]:
    """Execute alignment step and return filtered TSV for next step reuse."""
    with ALIGNMENT_LOG_PATH.open("w", encoding="utf-8") as alignment_log_handle:
        resource_plan = detect_alignment_resources(workers, bwa_threads_per_job)
        print_and_log(
            f"Alignment resource detection: CPU={resource_plan.cpu_count}, CPU load={resource_plan.cpu_usage_percent:.1f}%, "
            f"total memory={resource_plan.total_memory_gb:.2f} GB, available memory={resource_plan.available_memory_gb:.2f} GB, "
            f"workers={resource_plan.worker_count}, bwa threads per task={resource_plan.bwa_threads_per_job}.",
            alignment_log_handle,
        )

        reference_lengths = load_reference_lengths(reference_fasta.with_suffix(reference_fasta.suffix + ".fai"))
        print_and_log(
            f"Reference length statistics: KP={reference_lengths['KP_total_bases']} bp, EC={reference_lengths['EC_total_bases']} bp.",
            alignment_log_handle,
        )

        barcode_pairs = collect_barcode_pairs(barcode_fastq_dir, limit)
        if not barcode_pairs:
            raise RuntimeError("[Error] No barcode FASTQ pairs collected.")

        print_and_log(
            f"Collected {len(barcode_pairs)} barcodes, starting reads alignment. "
            f"Statistics criteria: MAPQ>={mapq_threshold}, identity>={strict_min_identity:.2f}, "
            f"aligned_fraction>={strict_min_aligned_fraction:.2f}, min_total_reads={min_total_reads}, "
            f"save_barcode_bam={save_barcode_bam}.",
            alignment_log_handle,
        )
        print_and_log(f"[Step 1/2] Alignment started: barcode count = {len(barcode_pairs)}.", pipeline_log_handle)

        task_inputs = [
            (
                barcode_record,
                str(reference_fasta),
                resource_plan.bwa_threads_per_job,
                reference_lengths["KP_total_bases"],
                reference_lengths["EC_total_bases"],
                mapq_threshold,
                strict_min_identity,
                strict_min_aligned_fraction,
                min_total_reads,
                str(BARCODE_RESULTS_DIR),
                save_barcode_bam,
            )
            for barcode_record in barcode_pairs
        ]

        result_rows: List[Dict[str, object]] = []
        error_rows: List[Dict[str, object]] = []

        with mp.Pool(processes=resource_plan.worker_count) as pool:
            for index, worker_output in enumerate(
                pool.imap_unordered(process_single_barcode, task_inputs, chunksize=10),
                start=1,
            ):
                if worker_output["status"] == "ok":
                    result_rows.append(worker_output["result"])
                else:
                    error_rows.append(
                        {
                            "sample_prefix": worker_output.get("sample_prefix", "NA"),
                            "error_message": worker_output.get("error_message", "Unknown error"),
                            "traceback": worker_output.get("traceback", ""),
                        }
                    )

                if index % 200 == 0 or index == len(task_inputs):
                    print_and_log(
                        f"Progress: completed {index}/{len(task_inputs)} barcodes; "
                        f"success {len(result_rows)}, failed {len(error_rows)}.",
                        alignment_log_handle,
                    )
                    sorted_progress_rows = sorted(
                        result_rows,
                        key=lambda row: (int(row["barcode_num"]), str(row["barcode_id"])),
                        reverse=True,
                    )
                    write_table_with_columns(sorted_progress_rows, ALIGNMENT_PROGRESS_TSV, ALIGNMENT_OUTPUT_COLUMNS, "\t")

        sorted_result_rows = sorted(
            result_rows,
            key=lambda row: (int(row["barcode_num"]), str(row["barcode_id"])),
            reverse=True,
        )
        filtered_result_rows = [
            row for row in sorted_result_rows if str(row["passed_min_total_reads_filter"]).lower() == "yes"
        ]

        write_table_with_columns(sorted_result_rows, ALIGNMENT_MAIN_TSV, ALIGNMENT_OUTPUT_COLUMNS, "\t")
        write_table_with_columns(sorted_result_rows, ALIGNMENT_MAIN_CSV, ALIGNMENT_OUTPUT_COLUMNS, ",")
        write_alignment_xlsx(sorted_result_rows, ALIGNMENT_MAIN_XLSX)
        write_table_with_columns(filtered_result_rows, ALIGNMENT_FILTERED_TSV, ALIGNMENT_OUTPUT_COLUMNS, "\t")
        write_table_with_columns(filtered_result_rows, ALIGNMENT_FILTERED_CSV, ALIGNMENT_OUTPUT_COLUMNS, ",")
        write_alignment_xlsx(filtered_result_rows, ALIGNMENT_FILTERED_XLSX)

        if error_rows:
            with ALIGNMENT_ERROR_TSV.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=["sample_prefix", "error_message", "traceback"], delimiter="\t")
                writer.writeheader()
                for row in error_rows:
                    writer.writerow(row)

        write_alignment_summary(
            barcode_fastq_dir=barcode_fastq_dir,
            reference_fasta=reference_fasta,
            sample_id=sample_id,
            total_barcodes=len(barcode_pairs),
            success_count=len(sorted_result_rows),
            filtered_count=len(filtered_result_rows),
            error_count=len(error_rows),
            reference_lengths=reference_lengths,
            resource_plan=resource_plan,
            rows_all=sorted_result_rows,
            rows_filtered=filtered_result_rows,
            mapq_threshold=mapq_threshold,
            strict_min_identity=strict_min_identity,
            strict_min_aligned_fraction=strict_min_aligned_fraction,
            min_total_reads=min_total_reads,
            save_barcode_bam=save_barcode_bam,
        )

        print_and_log(f"Full results TSV: {ALIGNMENT_MAIN_TSV}", alignment_log_handle)
        print_and_log(f"Filtered results TSV: {ALIGNMENT_FILTERED_TSV}", alignment_log_handle)
        print_and_log(f"Barcode independent result directory: {BARCODE_RESULTS_DIR}", alignment_log_handle)
        print_and_log("Alignment step completed.", alignment_log_handle)
        print_and_log(
            f"[Step 1/2] Alignment completed: success {len(sorted_result_rows)}, filtered retained {len(filtered_result_rows)}, failed {len(error_rows)}.",
            pipeline_log_handle,
        )

        return {
            "filtered_tsv": ALIGNMENT_FILTERED_TSV,
            "barcode_count": len(barcode_pairs),
            "success_count": len(sorted_result_rows),
            "filtered_count": len(filtered_result_rows),
            "error_count": len(error_rows),
        }


def run_threshold_sweep_step(
    filtered_tsv_path: Path,
    sample_id: str,
    purity_threshold_percent: float,
    reads_thresholds: List[int],
    pipeline_log_handle,
) -> Dict[str, object]:
    """Execute threshold sweep step based on filtered TSV."""
    with THRESHOLD_SWEEP_LOG_PATH.open("w", encoding="utf-8") as threshold_log_handle:
        if not filtered_tsv_path.exists():
            raise FileNotFoundError(f"[Error] Threshold sweep input file not found: {filtered_tsv_path}")

        resource_status = {
            "cpu_count": float(psutil.cpu_count(logical=True) or 1),
            "cpu_usage_percent": float(psutil.cpu_percent(interval=1)),
            "available_memory_gb": float(psutil.virtual_memory().available / (1024**3)),
        }
        cleaned_thresholds = validate_reads_thresholds(list(reads_thresholds))
        purity_threshold = purity_threshold_percent / 100.0

        print_and_log(
            f"Threshold sweep resource detection: CPU={int(resource_status['cpu_count'])}, "
            f"current load={resource_status['cpu_usage_percent']:.1f}%, "
            f"available memory={resource_status['available_memory_gb']:.2f} GB.",
            threshold_log_handle,
        )
        print_and_log(f"[Step 2/2] Reading input file: {filtered_tsv_path}", threshold_log_handle)
        print_and_log(
            f"[Step 2/2] Threshold sweep started: threshold count = {len(cleaned_thresholds)}, "
            f"purity threshold = {purity_threshold_percent:.0f}%.",
            pipeline_log_handle,
        )

        df = load_threshold_sweep_input_dataframe(filtered_tsv_path)
        print_and_log(f"Input barcode total count: {len(df)}", threshold_log_handle)

        summary_df = sweep_thresholds(df, cleaned_thresholds, purity_threshold)
        summary_df.to_csv(THRESHOLD_SWEEP_SUMMARY_TSV, sep="\t", index=False)
        summary_df.to_csv(THRESHOLD_SWEEP_SUMMARY_CSV, index=False)
        summary_df.to_excel(THRESHOLD_SWEEP_SUMMARY_XLSX, index=False)

        write_threshold_sweep_run_summary(
            input_file=filtered_tsv_path,
            sample_id=sample_id,
            purity_threshold_percent=purity_threshold_percent,
            reads_thresholds=cleaned_thresholds,
            input_barcode_count=len(df),
            resource_status=resource_status,
        )
        write_threshold_sweep_legend(
            input_file=filtered_tsv_path,
            input_barcode_count=len(df),
            purity_threshold_percent=purity_threshold_percent,
            reads_thresholds=cleaned_thresholds,
        )

        print_and_log(f"TSV results: {THRESHOLD_SWEEP_SUMMARY_TSV}", threshold_log_handle)
        print_and_log(f"CSV results: {THRESHOLD_SWEEP_SUMMARY_CSV}", threshold_log_handle)
        print_and_log(f"XLSX results: {THRESHOLD_SWEEP_SUMMARY_XLSX}", threshold_log_handle)
        print_and_log("Threshold sweep step completed.", threshold_log_handle)
        print_and_log(
            f"[Step 2/2] Threshold sweep completed: results written to {THRESHOLD_SWEEP_SUMMARY_TSV}.",
            pipeline_log_handle,
        )

        return {
            "input_barcode_count": len(df),
            "threshold_count": len(cleaned_thresholds),
            "summary_tsv": THRESHOLD_SWEEP_SUMMARY_TSV,
        }


def main() -> int:
    """Main pipeline entry: first align, then automatically run filtered threshold sweep."""
    args = parse_args()
    barcode_fastq_dir = Path(args.barcode_fastq_dir)
    reference_fasta = Path(args.reference_fasta)

    try:
        ensure_conda_environment()
        ensure_required_tools()

        purity_threshold_percent = float(args.purity_threshold_percent)
        if not (0 < purity_threshold_percent < 100):
            raise ValueError("[Error] purity_threshold_percent must be between 0 and 100.")

        version_number = detect_next_version(TASK_KEY)
        configure_output_paths(version_number, args.sample_id)
        ensure_master_paths(barcode_fastq_dir, reference_fasta)

        with PIPELINE_LOG_PATH.open("w", encoding="utf-8") as pipeline_log_handle:
            master_vm = psutil.virtual_memory()
            print_and_log("=" * 72, pipeline_log_handle)
            print_and_log("Align Barcodes Reads And Threshold Sweep", pipeline_log_handle)
            print_and_log("=" * 72, pipeline_log_handle)
            print_and_log(
                f"Master resource overview: CPU={psutil.cpu_count(logical=True) or 1}, "
                f"CPU load={psutil.cpu_percent(interval=1):.1f}%, "
                f"available memory={master_vm.available / (1024**3):.2f} GB.",
                pipeline_log_handle,
            )

            alignment_result = run_alignment_step(
                barcode_fastq_dir=barcode_fastq_dir,
                reference_fasta=reference_fasta,
                sample_id=args.sample_id,
                mapq_threshold=args.mapq_threshold,
                strict_min_identity=float(args.strict_min_identity),
                strict_min_aligned_fraction=float(args.strict_min_aligned_fraction),
                min_total_reads=args.min_total_reads,
                save_barcode_bam=args.save_barcode_bam,
                limit=args.limit,
                workers=args.workers,
                bwa_threads_per_job=args.bwa_threads_per_job,
                pipeline_log_handle=pipeline_log_handle,
            )

            threshold_result = run_threshold_sweep_step(
                filtered_tsv_path=Path(alignment_result["filtered_tsv"]),
                sample_id=args.sample_id,
                purity_threshold_percent=purity_threshold_percent,
                reads_thresholds=list(args.reads_thresholds),
                pipeline_log_handle=pipeline_log_handle,
            )

            backed_up_script = backup_current_script()
            append_readme_log(
                output_dir=OUTPUT_DIR,
                sample_id=args.sample_id,
                barcode_fastq_dir=barcode_fastq_dir,
                reference_fasta=reference_fasta,
                mapq_threshold=args.mapq_threshold,
                strict_min_identity=float(args.strict_min_identity),
                strict_min_aligned_fraction=float(args.strict_min_aligned_fraction),
                min_total_reads=args.min_total_reads,
                purity_threshold_percent=purity_threshold_percent,
                reads_thresholds=list(args.reads_thresholds),
            )

            print_and_log(f"Alignment subdirectory: {ALIGNMENT_DIR}", pipeline_log_handle)
            print_and_log(f"Reads_num_threshold_sweep subdirectory: {THRESHOLD_SWEEP_DIR}", pipeline_log_handle)
            print_and_log(
                f"Filtered barcode count: {alignment_result['filtered_count']}; "
                f"threshold count: {threshold_result['threshold_count']}.",
                pipeline_log_handle,
            )
            print_and_log(f"Script backup: {backed_up_script}", pipeline_log_handle)
            print_and_log(f"README log: {README_LOG}", pipeline_log_handle)
            print_and_log(f"Main output directory: {OUTPUT_DIR}", pipeline_log_handle)
            print_and_log("Integrated pipeline completed.", pipeline_log_handle)

        return 0
    except Exception as exc:
        print(f"[Error] Integrated pipeline failed: {exc}", file=sys.stderr)
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    sys.exit(main())
